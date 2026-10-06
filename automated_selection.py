"""Find NSE stocks with RSI above 60, positive MACD histogram, and no sideways regime."""

from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

import sideways
from scan_nse_all_stocks import calculate_macd, calculate_rsi


PROJECT_ROOT = Path(__file__).resolve().parent
DEFAULT_INPUT = PROJECT_ROOT / "masterdata_excels" / "EQUITY_L.csv"
OUTPUT_DIR = PROJECT_ROOT / "daily_nse_scans_alerts"
RSI_MINIMUM = 60.0
DEFAULT_YEARS = 2.0
DEFAULT_MIN_DURATION = 1
SIDEWAYS_CONFIG = sideways.SidewaysConfig(min_duration=DEFAULT_MIN_DURATION)
OUTPUT_COLUMNS = [
	"Symbol",
	"Name of Company",
	"Date",
	"Close",
	"Daily RSI(14)",
	"MACD Histogram",
	"Histogram Positive",
	"Sideways",
]


def load_master_symbols(input_path: Path) -> pd.DataFrame:
	"""Load unique NSE symbols and company names from the equity master CSV."""
	master = pd.read_csv(input_path, encoding="utf-8-sig")
	if "SYMBOL" not in master.columns:
		raise ValueError("EQUITY_L.csv must contain a SYMBOL column")
	if "NAME OF COMPANY" not in master.columns:
		master["NAME OF COMPANY"] = ""

	master = master[["SYMBOL", "NAME OF COMPANY"]].dropna(subset=["SYMBOL"]).copy()
	master["SYMBOL"] = master["SYMBOL"].astype(str).str.strip().str.upper()
	master["NAME OF COMPANY"] = (
		master["NAME OF COMPANY"].fillna("").astype(str).str.strip()
	)
	master = master[master["SYMBOL"] != ""].drop_duplicates(subset=["SYMBOL"])
	return master.reset_index(drop=True)


def evaluate_history(
	symbol: str,
	company_name: str,
	history: pd.DataFrame,
	config: sideways.SidewaysConfig = SIDEWAYS_CONFIG,
) -> tuple[dict[str, str | float | bool] | None, str]:
	"""Evaluate the latest completed daily bar against all three scan conditions."""
	if len(history) < config.effective_warmup:
		return None, (
			f"SKIP: only {len(history)} daily bars; need at least "
			f"{config.effective_warmup}"
		)

	close = pd.to_numeric(history["close"], errors="coerce").dropna()
	if close.empty:
		return None, "SKIP: no valid daily closing prices"

	daily_rsi = float(calculate_rsi(close, period=14).iloc[-1])
	macd_histogram = float(calculate_macd(close)["histogram"].iloc[-1])
	sideways_result = sideways.detect_sideways(history, config)
	is_sideways = bool(sideways_result["sideways"].iloc[-1])

	failures = []
	if pd.isna(daily_rsi) or daily_rsi <= RSI_MINIMUM:
		failures.append(f"daily RSI={daily_rsi:.2f} is not above {RSI_MINIMUM:.0f}")
	if pd.isna(macd_histogram) or macd_histogram <= 0:
		failures.append(f"MACD histogram={macd_histogram:.4f} is not positive")
	if is_sideways:
		failures.append("stock is sideways")
	result = {
		"Symbol": symbol,
		"Name of Company": company_name,
		"Date": history.index[-1].strftime("%Y-%m-%d"),
		"Close": round(float(close.iloc[-1]), 2),
		"Daily RSI(14)": round(daily_rsi, 2),
		"MACD Histogram": round(macd_histogram, 4),
		"Histogram Positive": bool(macd_histogram > 0),
		"Sideways": is_sideways,
	}
	if failures:
		return result, "FAIL: " + "; ".join(failures)

	return result, "PASS"


def evaluate_symbol(
	symbol: str,
	company_name: str,
	years: float,
	config: sideways.SidewaysConfig = SIDEWAYS_CONFIG,
) -> tuple[dict[str, str | float | bool] | None, str]:
	"""Fetch completed daily bars and evaluate one NSE symbol."""
	history = sideways.fetch_nse_data(symbol, years=years)
	return evaluate_history(symbol, company_name, history, config)


def scan_symbols(
	master: pd.DataFrame,
	years: float,
	limit: int = 0,
	config: sideways.SidewaysConfig = SIDEWAYS_CONFIG,
) -> list[dict[str, str | float | bool]]:
	"""Evaluate symbols from the master CSV and report each scan result."""
	if limit > 0:
		master = master.head(limit)

	matches = []
	total = len(master)
	for position, row in master.iterrows():
		symbol = row["SYMBOL"]
		company_name = row["NAME OF COMPANY"]
		try:
			result, status = evaluate_symbol(symbol, company_name, years, config)
		except Exception as error:
			result, status = None, f"SKIP: {error}"

		print(f"[{position + 1}/{total}] {symbol}: {status}")
		if result is not None and status == "PASS":
			matches.append(result)

	matches.sort(key=lambda match: float(match["Daily RSI(14)"]), reverse=True)
	return matches


def save_matches(
	matches: list[dict[str, str | float | bool]], output_path: Path
) -> None:
	"""Write qualifying symbols to a readable Excel workbook."""
	output_path.parent.mkdir(parents=True, exist_ok=True)
	pd.DataFrame(matches, columns=OUTPUT_COLUMNS).to_excel(output_path, index=False)

	workbook = load_workbook(output_path)
	worksheet = workbook.active
	worksheet.freeze_panes = "A2"
	worksheet.auto_filter.ref = worksheet.dimensions
	for cell in worksheet[1]:
		cell.font = Font(bold=True, color="FFFFFF")
		cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
	for column_cells in worksheet.columns:
		width = max(len(str(cell.value or "")) for cell in column_cells)
		worksheet.column_dimensions[column_cells[0].column_letter].width = min(width + 2, 40)
	workbook.save(output_path)


def main() -> None:
	parser = argparse.ArgumentParser(
		description="Scan NSE stocks for daily RSI, positive MACD histogram, and sideways status."
	)
	parser.add_argument(
		"--input", type=Path, default=DEFAULT_INPUT, help="Path to EQUITY_L.csv"
	)
	parser.add_argument(
		"--years", type=float, default=DEFAULT_YEARS,
		help="Years of daily history to fetch (default: 2)",
	)
	parser.add_argument(
		"--limit", type=int, default=0,
		help="Optional symbol limit for testing (0 = all)",
	)
	parser.add_argument(
		"--scan-date",
		type=date.fromisoformat,
		default=None,
		help="Output filename date (YYYY-MM-DD); skips the date prompt",
	)
	parser.add_argument(
		"--min-duration",
		type=int,
		default=DEFAULT_MIN_DURATION,
		help="Sideways bars required to confirm a range (default: 1)",
	)
	args = parser.parse_args()
	if args.years <= 0:
		parser.error("--years must be greater than 0")
	if args.limit < 0:
		parser.error("--limit cannot be negative")
	if args.min_duration < 1:
		parser.error("--min-duration must be at least 1")

	scan_date = args.scan_date
	if scan_date is None:
		today = date.today()
		date_text = input(
			f"Output date (YYYY-MM-DD; press Enter for {today}): "
		).strip()
		if date_text:
			try:
				scan_date = date.fromisoformat(date_text)
			except ValueError:
				parser.error("output date must use YYYY-MM-DD format")
		else:
			scan_date = today

	master = load_master_symbols(args.input)
	config = sideways.SidewaysConfig(min_duration=args.min_duration)
	matches = scan_symbols(master, args.years, args.limit, config)
	output_path = OUTPUT_DIR / f"automated_alert_{scan_date:%Y-%m-%d}.xlsx"
	save_matches(matches, output_path)
	print(f"Saved {len(matches)} qualifying stocks to {output_path}")


if __name__ == "__main__":
	main()
