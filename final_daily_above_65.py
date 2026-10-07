"""Find NSE stocks with strong daily RSI, green Hull trend, and positive MACD."""

from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path

import pandas as pd
import yfinance as yf
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from hull_trend_strategy import HullConfig, hull_trend_strategy
from scan_nse_all_stocks import calculate_macd, calculate_rsi
from stock_indicators import calculate_volume_signal


PROJECT_ROOT = Path(__file__).resolve().parent
DEFAULT_INPUT = PROJECT_ROOT / "masterdata_excels" / "EQUITY_L.csv"
OUTPUT_DIR = PROJECT_ROOT / "daily_nse_scans_alerts"
RSI_MINIMUM = 65.0
HISTORY_PERIOD = "5y"
OUTPUT_COLUMNS = [
	"Symbol",
	"Name of Company",
	"Date",
	"Close",
	"Daily RSI(14)",
	"Hull Trend",
	"MACD Signal",
	"Volume Signal",
]
MACD_SIGNAL_FILLS = {
	"Green": "C6EFCE",
	"Red": "FFC7CE",
}
VOLUME_SIGNAL_FILLS = {
	"Blue": "D9EAF7",
	"Green": "C6EFCE",
	"Red": "FFC7CE",
	"Grey (Noise)": "D9D9D9",
}


def load_master_symbols(input_path: Path) -> pd.DataFrame:
	"""Load unique NSE symbols and company names from the master CSV."""
	master = pd.read_csv(input_path, encoding="utf-8-sig")
	if "SYMBOL" not in master.columns:
		raise ValueError("EQUITY_L.csv must contain a SYMBOL column")

	if "NAME OF COMPANY" not in master.columns:
		master["NAME OF COMPANY"] = ""
	master = master[["SYMBOL", "NAME OF COMPANY"]].dropna(subset=["SYMBOL"]).copy()
	master["SYMBOL"] = master["SYMBOL"].astype(str).str.strip().str.upper()
	master["NAME OF COMPANY"] = master["NAME OF COMPANY"].fillna("").astype(str).str.strip()
	master = master[master["SYMBOL"] != ""].drop_duplicates(subset=["SYMBOL"])
	return master.reset_index(drop=True)


def fetch_daily_history(symbol: str, period: str) -> pd.DataFrame:
	"""Fetch daily OHLCV bars for one NSE symbol."""
	ticker = symbol if symbol.endswith(".NS") else f"{symbol}.NS"
	history = yf.Ticker(ticker).history(period=period, auto_adjust=False)
	if isinstance(history.columns, pd.MultiIndex):
		history.columns = history.columns.get_level_values(0)
	if history.empty or "Close" not in history.columns:
		raise ValueError("no daily closing-price history returned")
	history["Close"] = pd.to_numeric(history["Close"], errors="coerce")
	history = history.dropna(subset=["Close"])
	if history.empty:
		raise ValueError("no valid daily closing prices returned")
	return history


def evaluate_symbol(
	symbol: str, company_name: str, period: str
) -> tuple[dict[str, str | float | None] | None, str]:
	"""Return a qualifying output row and a concise PASS/FAIL explanation."""
	history = fetch_daily_history(symbol, period)
	if len(history) < 60:
		return None, f"FAIL: only {len(history)} daily bars; need at least 60"

	close = history["Close"]
	daily_rsi = float(calculate_rsi(close).iloc[-1])
	hull = hull_trend_strategy(history, HullConfig(length=24))
	hull_trend = str(hull["trend"].iloc[-1]).strip().lower()
	macd_histogram = float(calculate_macd(close)["histogram"].iloc[-1])
	latest_close = float(close.iloc[-1])
	volume_signal = calculate_volume_signal(history)

	failures = []
	if pd.isna(daily_rsi) or daily_rsi < RSI_MINIMUM:
		failures.append(f"daily RSI={daily_rsi:.2f} < {RSI_MINIMUM:.0f}")
	if hull_trend != "green":
		failures.append(f"Hull trend={hull_trend}")
	if pd.isna(macd_histogram) or macd_histogram <= 0:
		failures.append(f"MACD histogram={macd_histogram:.4f} is not positive")

	if failures:
		return None, "FAIL: " + "; ".join(failures)

	return {
		"Symbol": symbol.removesuffix(".NS"),
		"Name of Company": company_name,
		"Date": history.index[-1].strftime("%Y-%m-%d"),
		"Close": round(latest_close, 2),
		"Daily RSI(14)": round(daily_rsi, 2),
		"Hull Trend": "Green",
		"MACD Signal": "Green" if macd_histogram > 0 else "Red",
		"Volume Signal": volume_signal if volume_signal is not None else "None",
	}, "PASS"


def scan_symbols(
	master: pd.DataFrame, period: str, limit: int = 0
) -> list[dict[str, str | float | None]]:
	"""Evaluate symbols and print a PASS or FAIL result for every stock."""
	if limit > 0:
		master = master.head(limit)

	matches = []
	total = len(master)
	for position, row in master.iterrows():
		symbol = row["SYMBOL"]
		company_name = row["NAME OF COMPANY"]
		try:
			result, status = evaluate_symbol(symbol, company_name, period)
		except Exception as error:
			result, status = None, f"FAIL: {error}"

		print(f"[{position + 1}/{total}] {symbol}: {status}")
		if result is not None:
			matches.append(result)
	return matches


def save_matches(
	matches: list[dict[str, str | float | None]], output_path: Path
) -> None:
	"""Save qualifying stocks to an Excel workbook, including an empty result."""
	output_path.parent.mkdir(parents=True, exist_ok=True)
	pd.DataFrame(matches, columns=OUTPUT_COLUMNS).to_excel(output_path, index=False)
	workbook = load_workbook(output_path)
	worksheet = workbook.active
	macd_column = OUTPUT_COLUMNS.index("MACD Signal") + 1
	volume_column = OUTPUT_COLUMNS.index("Volume Signal") + 1
	for row_number in range(2, worksheet.max_row + 1):
		macd_cell = worksheet.cell(row=row_number, column=macd_column)
		macd_fill = MACD_SIGNAL_FILLS.get(macd_cell.value)
		if macd_fill:
			macd_cell.fill = PatternFill(fill_type="solid", fgColor=macd_fill)

		signal_cell = worksheet.cell(row=row_number, column=volume_column)
		fill_color = VOLUME_SIGNAL_FILLS.get(signal_cell.value)
		if fill_color:
			signal_cell.fill = PatternFill(fill_type="solid", fgColor=fill_color)
	workbook.save(output_path)


def main() -> None:
	parser = argparse.ArgumentParser(
		description="Scan daily NSE RSI, Hull trend, and MACD conditions."
	)
	parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="Path to EQUITY_L.csv")
	parser.add_argument("--period", default=HISTORY_PERIOD, help="Yahoo Finance daily history period")
	parser.add_argument("--limit", type=int, default=0, help="Optional symbol limit for testing (0 = all)")
	parser.add_argument("--scan-date", type=date.fromisoformat, default=date.today(), help="Date for the output filename (YYYY-MM-DD)")
	args = parser.parse_args()

	master = load_master_symbols(args.input)
	matches = scan_symbols(master, args.period, args.limit)
	output_path = OUTPUT_DIR / f"daily_green_rsi_{args.scan_date:%Y-%m-%d}.xlsx"
	save_matches(matches, output_path)
	print(f"Saved {len(matches)} qualifying stocks to {output_path}")


if __name__ == "__main__":
	main()
