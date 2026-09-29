"""
auto_invest_suggestions.py

Reads a screener CSV (columns: Sr., Stock Name, Symbol, close, %_change,
volume, sector, industry), pulls weekly NSE data for each Symbol, computes:

    - Hull Trend Strategy signal (green/red)   [hull_trend_strategy.py]
    - Weekly RSI(14), Wilder's RMA             (matches Chartink)
    - Weekly EMA10 / EMA20 / EMA50

...and writes an .xlsx with only the stocks that satisfy:

    (weekly RSI >= 55  AND  Hull trend == green  AND  Close > EMA50)
    OR
    (weekly RSI >= 65)

Output columns: Symbol, Stock Name, Sector, Industry, Trend, Weekly RSI(14),
Weekly EMA10, Weekly EMA20, Weekly EMA50.

Depends on hull_trend_strategy.py (from the earlier conversation) being in
the same folder — this script imports hull_ma / triangular_hull_ma / the
crossover logic from there rather than re-implementing it, so the Hull
signal always stays identical to that script.

Usage:
    python3 auto_invest_suggestions.py screener.csv
    python3 auto_invest_suggestions.py screener.csv --output suggestions.xlsx
    python3 auto_invest_suggestions.py screener.csv --period 5y --hull-length 24
"""

from __future__ import annotations

import argparse
import sys
import time
from dataclasses import dataclass

import numpy as np
import pandas as pd
import yfinance as yf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from hull_trend_strategy import HullConfig, hull_trend_strategy

REQUIRED_CSV_COLS = [
    "Sr.", "Stock Name", "Symbol", "close", "%_change", "volume", "sector", "industry",
]

FONT_NAME = "Arial"
HEADER_FILL = "1F4E78"
GREEN_FILL = "C6EFCE"
RED_FILL = "FFC7CE"


# --------------------------------------------------------------------------
# Indicators
# --------------------------------------------------------------------------

def rsi_wilder(close: pd.Series, length: int = 14) -> pd.Series:
    """
    Wilder's RMA-based RSI (matches Chartink), via ewm(alpha=1/length,
    adjust=False) — same approach already established in the NSE screening
    pipeline for daily RSI, applied here to weekly bars.
    """
    delta = close.diff()
    gain = delta.clip(lower=0)
    loss = -delta.clip(upper=0)
    avg_gain = gain.ewm(alpha=1 / length, min_periods=length, adjust=False).mean()
    avg_loss = loss.ewm(alpha=1 / length, min_periods=length, adjust=False).mean()
    rs = avg_gain / avg_loss
    rsi = 100 - (100 / (1 + rs))
    rsi = rsi.where(avg_loss != 0, 100.0)  # no losses in window -> RSI 100
    return rsi


def ema(close: pd.Series, span: int) -> pd.Series:
    return close.ewm(span=span, adjust=False).mean()


# --------------------------------------------------------------------------
# Per-symbol evaluation
# --------------------------------------------------------------------------

@dataclass
class ScreenConfig:
    hull_length: int = 24
    rsi_length: int = 14
    period: str = "5y"      # weekly history window pulled per symbol
    rsi_soft_floor: float = 55.0   # needs green trend + close > EMA50 too
    rsi_hard_floor: float = 65.0   # qualifies on RSI alone
    ema_trend_len: int = 50
    pause_sec: float = 0.3  # small delay between requests, be polite to yfinance


def fetch_weekly_symbol(symbol: str, period: str) -> pd.DataFrame | None:
    ticker = symbol if symbol.upper().endswith(".NS") else f"{symbol.upper()}.NS"
    try:
        df = yf.download(
            ticker,
            period=period,
            interval="1wk",
            auto_adjust=False,
            progress=False,
        )
    except Exception as exc:  # network / ticker errors
        print(f"  ! {symbol}: fetch failed ({exc})")
        return None

    if df.empty:
        print(f"  ! {symbol}: no data returned")
        return None
    if isinstance(df.columns, pd.MultiIndex):
        df.columns = df.columns.get_level_values(0)
    return df


def evaluate_symbol(symbol: str, cfg: ScreenConfig) -> dict | None:
    """Returns a dict of computed indicator values for one symbol, or None
    if data was unavailable / insufficient."""
    df = fetch_weekly_symbol(symbol, cfg.period)
    if df is None:
        return None

    min_bars = max(cfg.hull_length * 2, cfg.ema_trend_len * 2, 60)
    if len(df) < min_bars:
        print(f"  ! {symbol}: only {len(df)} weekly bars, need >= {min_bars} — skipping")
        return None

    hull = hull_trend_strategy(df, HullConfig(length=cfg.hull_length))
    rsi = rsi_wilder(df["Close"], cfg.rsi_length)
    ema10 = ema(df["Close"], 10)
    ema20 = ema(df["Close"], 20)
    ema50 = ema(df["Close"], cfg.ema_trend_len)

    last = -1
    trend = hull["trend"].iloc[last]
    close = df["Close"].iloc[last]
    rsi_val = rsi.iloc[last]
    ema10_val = ema10.iloc[last]
    ema20_val = ema20.iloc[last]
    ema50_val = ema50.iloc[last]

    if pd.isna(rsi_val) or pd.isna(ema50_val) or pd.isna(trend):
        print(f"  ! {symbol}: indicators not fully warmed up — skipping")
        return None

    return {
        "trend": trend,
        "close": float(close),
        "rsi": float(rsi_val),
        "ema10": float(ema10_val),
        "ema20": float(ema20_val),
        "ema50": float(ema50_val),
    }


def passes_filter(vals: dict, cfg: ScreenConfig) -> bool:
    soft = (
        vals["rsi"] >= cfg.rsi_soft_floor
        and vals["trend"] == "green"
        and vals["close"] > vals["ema50"]
    )
    hard = vals["rsi"] >= cfg.rsi_hard_floor
    return soft or hard


# --------------------------------------------------------------------------
# Main screening run
# --------------------------------------------------------------------------

def run_screen(csv_path: str, cfg: ScreenConfig) -> pd.DataFrame:
    universe = pd.read_csv(csv_path)
    missing = [c for c in REQUIRED_CSV_COLS if c not in universe.columns]
    if missing:
        raise ValueError(f"Input CSV is missing required columns: {missing}")

    rows = []
    total = len(universe)
    for i, row in universe.iterrows():
        symbol = str(row["Symbol"]).strip()
        stock_name = row["Stock Name"]
        sector = row["sector"]
        industry = row["industry"]

        print(f"[{i + 1}/{total}] {symbol} ...")
        vals = evaluate_symbol(symbol, cfg)
        time.sleep(cfg.pause_sec)

        if vals is None:
            continue

        if passes_filter(vals, cfg):
            rows.append({
                "Symbol": symbol,
                "Stock Name": stock_name,
                "Sector": sector,
                "Industry": industry,
                "Trend": vals["trend"].capitalize(),
                "Weekly RSI(14)": round(vals["rsi"], 2),
                "Weekly EMA10": round(vals["ema10"], 2),
                "Weekly EMA20": round(vals["ema20"], 2),
                "Weekly EMA50": round(vals["ema50"], 2),
            })
            print(f"  -> QUALIFIES  (trend={vals['trend']}, rsi={vals['rsi']:.2f})")

    return pd.DataFrame(rows)


# --------------------------------------------------------------------------
# Excel output
# --------------------------------------------------------------------------

def write_xlsx(df: pd.DataFrame, output_path: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Suggestions"

    headers = [
        "Symbol", "Stock Name", "Sector", "Industry",
        "Trend", "Weekly RSI(14)", "Weekly EMA10", "Weekly EMA20", "Weekly EMA50",
    ]
    ws.append(headers)
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = Font(name=FONT_NAME, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for _, row in df.iterrows():
        ws.append([
            row["Symbol"], row["Stock Name"], row["Sector"], row["Industry"],
            row["Trend"], row["Weekly RSI(14)"], row["Weekly EMA10"],
            row["Weekly EMA20"], row["Weekly EMA50"],
        ])

    for r in range(2, ws.max_row + 1):
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).font = Font(name=FONT_NAME)
        trend_cell = ws.cell(row=r, column=5)
        if trend_cell.value == "Green":
            trend_cell.fill = PatternFill("solid", fgColor=GREEN_FILL)
        elif trend_cell.value == "Red":
            trend_cell.fill = PatternFill("solid", fgColor=RED_FILL)

    widths = [12, 28, 20, 24, 9, 15, 13, 13, 13]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"

    wb.save(output_path)


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Auto-invest suggestions screener")
    parser.add_argument("csv_path", help="Input screener CSV path")
    parser.add_argument("--output", default="auto_invest_suggestions.xlsx", help="Output .xlsx path")
    parser.add_argument("--period", default="5y", help="yfinance weekly history window (e.g. 3y, 5y, max)")
    parser.add_argument("--hull-length", type=int, default=24, help="Hull Trend Strategy length")
    args = parser.parse_args()

    cfg = ScreenConfig(hull_length=args.hull_length, period=args.period)

    print(f"Screening {args.csv_path} ...\n")
    result = run_screen(args.csv_path, cfg)

    if result.empty:
        print("\nNo stocks passed the filter — no output file written.")
        sys.exit(0)

    result = result.sort_values("Weekly RSI(14)", ascending=False).reset_index(drop=True)
    write_xlsx(result, args.output)
    print(f"\n{len(result)} stocks qualified. Written to {args.output}")
