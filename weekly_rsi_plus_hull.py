from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
import pandas as pd
import yfinance as yf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from hull_trend_strategy import HullConfig, hull_trend_strategy


DEFAULT_INPUT = Path(__file__).resolve().parent / "masterdata_excels" / "EQUITY_L.csv"
DEFAULT_OUTPUT = Path(__file__).resolve().parent / "daily_nse_scans_alerts" / "weekly_green_rsi_above_50.xlsx"


def ema(close: pd.Series, span: int) -> pd.Series:
    return close.ewm(span=span, adjust=False).mean()


def rsi_wilder(close: pd.Series, length: int = 14) -> pd.Series:
    """Wilder RSI (matching Chartink-style weekly RSI)."""
    delta = close.diff()
    gain = delta.clip(lower=0)
    loss = -delta.clip(upper=0)

    avg_gain = gain.ewm(alpha=1 / length, min_periods=length, adjust=False).mean()
    avg_loss = loss.ewm(alpha=1 / length, min_periods=length, adjust=False).mean()
    rs = avg_gain / avg_loss.replace(0, np.nan)
    rsi = 100 - (100 / (1 + rs))
    return rsi.where(avg_loss != 0, 100.0)


def fetch_weekly(symbol: str, period: str = "5y") -> pd.DataFrame | None:
    ticker = symbol if symbol.upper().endswith(".NS") else f"{symbol.upper()}.NS"
    try:
        df = yf.download(
            ticker,
            period=period,
            interval="1wk",
            auto_adjust=False,
            progress=False,
        )
    except Exception:
        return None

    if df.empty:
        return None
    if isinstance(df.columns, pd.MultiIndex):
        df.columns = df.columns.get_level_values(0)
    df.index.name = "Date"
    return df


def fetch_daily(symbol: str, period: str = "6mo") -> pd.DataFrame | None:
    ticker = symbol if symbol.upper().endswith(".NS") else f"{symbol.upper()}.NS"
    try:
        df = yf.download(
            ticker,
            period=period,
            interval="1d",
            auto_adjust=False,
            progress=False,
        )
    except Exception:
        return None

    if df.empty:
        return None
    if isinstance(df.columns, pd.MultiIndex):
        df.columns = df.columns.get_level_values(0)
    df.index.name = "Date"
    return df


def load_master_symbols(input_path: Path) -> pd.DataFrame:
    df = pd.read_csv(input_path)
    required = {"SYMBOL", "NAME OF COMPANY"}
    missing = required - set(df.columns)
    if missing:
        raise ValueError(f"CSV is missing required columns: {sorted(missing)}")

    cleaned = df[["SYMBOL", "NAME OF COMPANY"]].copy()
    cleaned = cleaned.dropna(subset=["SYMBOL", "NAME OF COMPANY"]).copy()
    cleaned["SYMBOL"] = cleaned["SYMBOL"].astype(str).str.strip().str.upper()
    cleaned["NAME OF COMPANY"] = cleaned["NAME OF COMPANY"].astype(str).str.strip()
    cleaned = cleaned[cleaned["SYMBOL"] != ""]
    cleaned = cleaned.drop_duplicates(subset=["SYMBOL"], keep="first")
    return cleaned.reset_index(drop=True)


def evaluate_symbol(symbol: str, company_name: str, period: str = "5y") -> dict | None:
    df = fetch_weekly(symbol, period=period)
    if df is None or len(df) < 60:
        return None

    hull = hull_trend_strategy(df, HullConfig(length=24))
    weekly_rsi = rsi_wilder(df["Close"], length=14)
    weekly_ema50 = ema(df["Close"], 50)

    if len(hull) == 0 or len(weekly_rsi) == 0 or len(weekly_ema50) == 0:
        return None

    trend = str(hull["trend"].iloc[-1]).strip().lower()
    weekly_rsi_value = float(weekly_rsi.iloc[-1])
    close_value = float(df["Close"].iloc[-1])
    ema50_value = float(weekly_ema50.iloc[-1])

    if pd.isna(weekly_rsi_value) or pd.isna(ema50_value):
        return None

    daily_df = fetch_daily(symbol, period="6mo")
    daily_rsi_value = None
    if daily_df is not None and len(daily_df) >= 30:
        daily_rsi = rsi_wilder(daily_df["Close"], length=14)
        daily_rsi_value = float(daily_rsi.iloc[-1])

    is_above_ema50 = close_value > ema50_value
    print(f"  -> Trend={trend}, Weekly RSI={weekly_rsi_value:.2f}, Close={close_value:.2f}, EMA50={ema50_value:.2f}, Daily RSI={daily_rsi_value if daily_rsi_value is not None else 'N/A'}")

    if trend == "green" and weekly_rsi_value > 50 and is_above_ema50:
        return {
            "Symbol": symbol,
            "Name of Company": company_name,
            "Trend": "Green",
            "Weekly RSI(14)": round(weekly_rsi_value, 2),
            "Weekly Close": round(close_value, 2),
            "Weekly EMA50": round(ema50_value, 2),
            "Daily RSI(14)": round(daily_rsi_value, 2) if daily_rsi_value is not None else None,
        }

    return None


def build_output_excel(rows: list[dict], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    df = pd.DataFrame(rows)
    if df.empty:
        df = pd.DataFrame(columns=["Symbol", "Name of Company", "Trend", "Weekly RSI(14)", "Weekly Close", "Weekly EMA50", "Daily RSI(14)"])

    headers = ["Symbol", "Name of Company", "Trend", "Weekly RSI(14)", "Weekly Close", "Weekly EMA50", "Daily RSI(14)"]
    sheet_ranges = [
        ("Weekly RSI Above 65", df[df["Weekly RSI(14)"] > 65]),
        ("Weekly RSI 55-65", df[(df["Weekly RSI(14)"] >= 55) & (df["Weekly RSI(14)"] <= 65)]),
        ("Weekly RSI 50-55", df[(df["Weekly RSI(14)"] > 50) & (df["Weekly RSI(14)"] < 55)]),
    ]

    wb = Workbook()
    for sheet_index, (sheet_name, sheet_df) in enumerate(sheet_ranges):
        ws = wb.active if sheet_index == 0 else wb.create_sheet()
        ws.title = sheet_name
        ws.append(headers)

        for col_idx, _ in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = Font(name="Arial", bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        sheet_df = sheet_df.sort_values("Weekly RSI(14)", ascending=False)
        for _, row in sheet_df.iterrows():
            ws.append([row[header] for header in headers])

        for row_idx in range(2, ws.max_row + 1):
            for col_idx in range(1, len(headers) + 1):
                ws.cell(row=row_idx, column=col_idx).font = Font(name="Arial")

        for col_idx, width in enumerate([12, 35, 10, 15, 12, 12, 12], start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:G{ws.max_row}"

    wb.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Filter NSE stocks by weekly Hull green trend + weekly RSI > 50.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="Path to EQUITY_L.csv")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT, help="Output Excel path")
    parser.add_argument("--period", default="5y", help="Yahoo Finance history period, e.g. 2y, 5y")
    parser.add_argument("--limit", type=int, default=0, help="Optional limit on number of symbols to process for testing (0 = all)")
    args = parser.parse_args()

    master = load_master_symbols(args.input)
    if args.limit and args.limit > 0:
        master = master.head(args.limit)

    results = []
    total = len(master)
    for idx, row in master.iterrows():
        symbol = str(row["SYMBOL"]).strip()
        company_name = str(row["NAME OF COMPANY"]).strip()
        print(f"[{idx + 1}/{total}] {symbol} ...")
        result = evaluate_symbol(symbol, company_name, period=args.period)
        if result:
            results.append(result)
            print(f"  -> PASS  RSI={result['Weekly RSI(14)']}  Trend={result['Trend']}")

    build_output_excel(results, args.output)
    print(f"\nSaved {len(results)} qualifying stocks to: {args.output}")


if __name__ == "__main__":
    main()
